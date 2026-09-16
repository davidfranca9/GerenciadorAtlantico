"""A Autorizacao de Carregamento sai no modelo que a Fertimaxi mandou.

O modelo e um cartao vertical - rotulo em B, valor em C - e o arquivo traz
dois blocos. Um pedido vira um bloco; os testes travam que o primeiro bloco
do modelo e copiado fielmente (rotulos, bordas, cores, mesclagens) e que
nada do segundo bloco de exemplo sobra quando ha um pedido so.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.servicos.documentos import gerar_autorizacao_xlsx  # noqa: E402

MODELO = Path(__file__).resolve().parents[1] / "dados" / "Autorizacao de Carregamento FERTIMAXI.xlsx"

ROTULOS = ["Cliente:", "Pedido:", "Quantidade:", "Produto:", "Embalagem:", "Transportador:",
           "Motorista:", "CPF:", "Modelo Veiculo:", "Placa:", "Telefone:"]

PEDIDO_WAGMAR = {
    "cliente": "WAGMAR JOSE DE OLIVEIRA", "contrato": "41556", "toneladas": "32",
    "produto": "UREIA PRILL MICROGRANULADA 46% N", "embalagem": "SACARIA", "cidade": "Montes Claros",
}


def gerar(tmp_path, produtos, **motorista):
    destino = tmp_path / "autorizacao.xlsx"
    base = dict(motorista="TALISSON JUNIOR GUIMARAES RIBEIRO", cpf="12159781622",
                telefone="(38) 99999-0000", placas=("PFJ2I64", "nzb4h89", ""), modelo_veiculo="SIDER")
    base.update(motorista)
    gerar_autorizacao_xlsx(str(MODELO), str(destino), produtos, **base)
    return load_workbook(destino).active


def valores(ws, topo):
    """{rotulo: valor} de um bloco que comeca na linha `topo`."""
    return {ws[f"B{topo + 3 + i}"].value: ws[f"C{topo + 3 + i}"].value for i in range(len(ROTULOS))}


def test_um_pedido_preenche_o_bloco_do_modelo(tmp_path):
    ws = gerar(tmp_path, [PEDIDO_WAGMAR])
    assert ws["B2"].value == "FERTIMAXI"
    assert ws["B3"].value == "Autorização De Carregamento"
    assert valores(ws, 2) == {
        "Cliente:": "WAGMAR JOSE DE OLIVEIRA",
        "Pedido:": 41556,
        "Quantidade:": "32 t",
        "Produto:": "UREIA PRILL MICROGRANULADA 46% N",
        "Embalagem:": "SACARIA",
        "Transportador:": "ATLÂNTICO FERTLOG",
        "Motorista:": "TALISSON JUNIOR GUIMARAES RIBEIRO",
        "CPF:": "121.597.816-22",
        "Modelo Veiculo:": "SIDER",
        "Placa:": "PFJ-2I64 / NZB-4H89",
        "Telefone:": "(38) 99999-0000",
    }


def test_um_pedido_nao_deixa_sobra_do_segundo_bloco_de_exemplo(tmp_path):
    ws = gerar(tmp_path, [PEDIDO_WAGMAR])
    assert all(ws[f"{c}{r}"].value is None for r in range(16, 31) for c in "BC")
    assert sorted(str(m) for m in ws.merged_cells.ranges) == ["B2:C2", "B3:C3"]
    assert ws.print_area == "'Planilha1'!$B$2:$C$15"


def test_cada_pedido_vira_um_bloco(tmp_path):
    outro = dict(PEDIDO_WAGMAR, cliente="ALYSSON SANTOS AGUIAR", contrato="41416", toneladas="4.5")
    terceiro = dict(PEDIDO_WAGMAR, contrato="40778", embalagem="BIG BAG")
    ws = gerar(tmp_path, [PEDIDO_WAGMAR, outro, terceiro])

    for topo in (2, 17, 32):
        assert ws[f"B{topo}"].value == "FERTIMAXI"
        assert [ws[f"B{topo + 3 + i}"].value for i in range(len(ROTULOS))] == ROTULOS
    assert valores(ws, 17)["Cliente:"] == "ALYSSON SANTOS AGUIAR"
    assert valores(ws, 17)["Quantidade:"] == "4,5 t"
    assert valores(ws, 32)["Pedido:"] == 40778
    assert valores(ws, 32)["Embalagem:"] == "BIG BAG"
    assert {"B17:C17", "B18:C18", "B32:C32", "B33:C33"} <= {str(m) for m in ws.merged_cells.ranges}
    assert ws.print_area == "'Planilha1'!$B$2:$C$45"


def test_o_estilo_do_modelo_e_copiado_pros_blocos_novos(tmp_path):
    ws = gerar(tmp_path, [PEDIDO_WAGMAR] * 3)
    modelo = load_workbook(MODELO).active
    for linha_modelo, linha_nova in ((3, 33), (5, 35), (15, 45)):
        for col in "BC":
            original, copia = modelo[f"{col}{linha_modelo}"], ws[f"{col}{linha_nova}"]
            assert copia.font.b == original.font.b and copia.font.sz == original.font.sz
            assert copia.border.left.style == original.border.left.style
            assert copia.border.bottom.style == original.border.bottom.style
            assert copia.fill.fgColor.rgb == original.fill.fgColor.rgb
        assert ws.row_dimensions[linha_nova].height == modelo.row_dimensions[linha_modelo].height


def test_dois_blocos_por_pagina_como_no_modelo(tmp_path):
    ws = gerar(tmp_path, [PEDIDO_WAGMAR] * 5)
    # Quebra antes do 3o e do 5o bloco (linhas 32 e 62).
    assert [b.id for b in ws.row_breaks.brk] == [31, 61]


def test_sem_pedido_ainda_sai_um_bloco_com_o_motorista(tmp_path):
    ws = gerar(tmp_path, [])
    assert valores(ws, 2)["Motorista:"] == "TALISSON JUNIOR GUIMARAES RIBEIRO"
    assert valores(ws, 2)["Cliente:"] in (None, "")


@pytest.mark.parametrize("escolhido, impresso", [
    ("GRANELEIRO", "GRANELEIRO"),
    ("grade baixa", "GRADE BAIXA"),
    ("SIDER", "SIDER"),
])
def test_modelo_do_veiculo_e_o_escolhido_na_ordem_de_coleta(tmp_path, escolhido, impresso):
    assert valores(gerar(tmp_path, [PEDIDO_WAGMAR], modelo_veiculo=escolhido), 2)["Modelo Veiculo:"] == impresso


def test_sem_modelo_escolhido_fica_em_branco_mesmo_com_placas(tmp_path):
    # Nao deduz pelas placas: cavalo + carreta pode ser qualquer carroceria.
    ws = gerar(tmp_path, [PEDIDO_WAGMAR], modelo_veiculo="", placas=("PFJ2I64", "NZB4H89", ""))
    assert (valores(ws, 2)["Modelo Veiculo:"] or "") == ""


def test_cpf_e_placa_sem_mascara_valida_saem_como_vieram(tmp_path):
    ws = gerar(tmp_path, [PEDIDO_WAGMAR], cpf="123", placas=("TESTE", "", ""))
    assert valores(ws, 2)["CPF:"] == "123"
    assert valores(ws, 2)["Placa:"] == "TESTE"
