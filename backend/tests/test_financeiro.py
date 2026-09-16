"""Financeiro: as contas das planilhas "Fluxo Caixa" e "Controle de
carregamentos", agora no sistema.

As planilhas daqui sao montadas no teste com a mesma estrutura das reais
(abas, linhas e formulas), com numeros inventados - os arquivos de verdade
tem dados pessoais e nao entram no repositorio.
"""
from __future__ import annotations

import io
import sys
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.auth import get_current_user, require_admin  # noqa: E402
from app.database import get_db  # noqa: E402
from app.models import (  # noqa: E402
    CarregamentoFinanceiro, ContaAvulsa, ContaBancaria, Despesa, Divida, LancamentoCaixa, MetaMensal, PagamentoAgenda,
)
from app.routers import financeiro as rotas  # noqa: E402
from app.servicos import financeiro as fin  # noqa: E402
from app.servicos import financeiro_importacao as imp  # noqa: E402
from tests.apoio_documentos import banco_em_memoria  # noqa: E402

TABELAS = (ContaBancaria, LancamentoCaixa, MetaMensal, CarregamentoFinanceiro, Despesa, ContaAvulsa, PagamentoAgenda, Divida)
DIA = date(2026, 9, 15)


@pytest.fixture
def db():
    sessao = banco_em_memoria(*TABELAS)
    yield sessao
    sessao.close()


def conta(db, nome="Nubank", saldo=1000.0, em=DIA):
    c = ContaBancaria(nome=nome, instituicao=nome, cor="#000", saldo_inicial=saldo, saldo_inicial_em=em)
    db.add(c)
    db.flush()
    return c


def lancar(db, c, dia, tipo, valor, descricao="x", forma="PIX"):
    return fin.criar_lancamento(db, conta_id=c.id, data=dia, tipo=tipo, valor=valor, descricao=descricao, forma=forma)


# --------------------------------------------------------------------------
# Caixa
# --------------------------------------------------------------------------


def test_saldo_anda_pra_frente_e_pra_tras_do_saldo_inicial(db):
    nu = conta(db, saldo=1000)
    lancar(db, nu, date(2026, 9, 14), "saida", 200)   # antes: ja estava no saldo inicial
    lancar(db, nu, DIA, "entrada", 500)
    lancar(db, nu, date(2026, 9, 16), "saida", 100)
    assert fin.saldo_no_inicio(db, nu, date(2026, 9, 14)) == 1200
    assert fin.saldo_no_inicio(db, nu, DIA) == 1000
    assert fin.saldo_no_inicio(db, nu, date(2026, 9, 16)) == 1500
    assert fin.saldo_no_inicio(db, nu, date(2026, 9, 17)) == 1400


def test_transferencia_entre_contas_nao_conta_como_receita_nem_gasto(db):
    nu, bb = conta(db, "Nubank", 1000), conta(db, "Banco do Brasil", 5000)
    fin.criar_transferencia(db, origem_id=bb.id, destino_id=nu.id, data=DIA, valor=3000)
    lancar(db, nu, DIA, "saida", 120, "Uber")
    resumo = fin.resumo_caixa(db, DIA, DIA)
    t = resumo["totais"]
    assert (t["abertura"], t["fechamento"]) == (6000, 5880)
    assert (t["entradas"], t["saidas"]) == (3000, 3120)
    assert (t["entradas_externas"], t["saidas_externas"], t["transferencias"], t["resultado"]) == (0, 120, 3000, -120)
    entrada = next(l for l in resumo["lancamentos"] if l["tipo"] == "entrada")
    assert entrada["contraparte"] == "Banco do Brasil"


def test_excluir_uma_ponta_da_transferencia_apaga_as_duas(db):
    nu, bb = conta(db, "Nubank"), conta(db, "Banco do Brasil")
    saida, _ = fin.criar_transferencia(db, origem_id=bb.id, destino_id=nu.id, data=DIA, valor=50)
    assert fin.excluir_lancamento(db, saida.id) == 2
    assert db.query(LancamentoCaixa).count() == 0


def test_serie_do_saldo_tem_um_ponto_por_dia(db):
    nu = conta(db, saldo=100, em=date(2026, 9, 1))
    lancar(db, nu, date(2026, 9, 10), "entrada", 50)
    serie = fin.resumo_caixa(db, DIA, DIA, dias_serie=30)["serie"]
    assert len(serie) == 30 and serie[-1] == {"data": "2026-09-15", "saldo": 150}
    assert next(p for p in serie if p["data"] == "2026-09-09")["saldo"] == 100


@pytest.mark.parametrize("texto, esperado", [
    ("PIX: Bianca Maria ", ("PIX", "Bianca Maria")),
    ("PIX:Uber ", ("PIX", "Uber")),
    ("TRNSF: Banco do Brasil", ("TRANSFERENCIA", "Banco do Brasil")),
    ("BOLETOS: Equilibrio", ("BOLETO", "Equilibrio")),
    ("DEBITO: Consorcio Sicredi", ("DEBITO", "Consorcio Sicredi")),
    ("RENDE FACIL", ("RENDIMENTO", "RENDE FACIL")),
    ("Aluguel", ("OUTRO", "Aluguel")),
])
def test_forma_sai_do_prefixo_da_planilha(texto, esperado):
    assert fin.separar_forma(texto) == esperado


# --------------------------------------------------------------------------
# Resultado do mes
# --------------------------------------------------------------------------


def carregamento(db, **campos):
    base = dict(competencia="2026-09", ctes="5003", motorista="EDILSON", fabrica="Fertimaxi", peso=32,
                frete_empresa_ton=270, frete_motorista_ton=210, agenciamento_ton=10, comissao_ton=0, contratante="Queiroz")
    base.update(campos)
    c = CarregamentoFinanceiro(**base)
    db.add(c)
    db.flush()
    return c


def test_linha_da_planilha_da_o_mesmo_liquido(db):
    # 5003: 32 t x (270 - 210 - 10) = 1.600
    assert fin.totais_carregamento(carregamento(db))["liquido"] == 1600


def test_total_fechado_manda_sobre_o_por_tonelada(db):
    # 5004: frete do motorista fechado em 9.000, sem valor por tonelada.
    c = carregamento(db, ctes="5004", frete_empresa_ton=340, frete_motorista_ton=None, frete_motorista_total=9000)
    totais = fin.totais_carregamento(c)
    assert (totais["frete_motorista"], totais["liquido"]) == (9000, 1560)


def test_cancelado_nao_entra_na_conta(db):
    carregamento(db)
    carregamento(db, ctes="5010", motorista="", cancelado=True)
    resumo = fin.resultado_mensal(db, "2026-09", hoje=date(2026, 9, 16))["resumo"]
    assert (resumo["carregamentos"], resumo["cancelados"], resumo["toneladas"], resumo["lucro_bruto"]) == (1, 1, 32, 1600)


def test_meta_precificacao_e_cascata(db):
    carregamento(db)
    db.add(MetaMensal(competencia="2026-09", meta_toneladas=3000))
    db.add(Despesa(escopo="empresa", grupo="Sistemas", descricao="Bsoft", valor=800, competencia_inicio="2026-09",
                   entra_precificacao=True))
    db.add(Despesa(escopo="empresa", grupo="Custos da operação", descricao="Seguro RCTRC", valor=800,
                   competencia_inicio="2026-09", conta_no_resultado=False, entra_precificacao=True))
    db.add(Despesa(escopo="pessoal", grupo="Casa", descricao="Aluguel", valor=500, competencia_inicio="2026-09"))
    db.flush()
    r = fin.resultado_mensal(db, "2026-09", hoje=date(2026, 9, 16))
    # 16 a 30/09 sem os domingos (20 e 27): 13 dias de carregamento.
    assert r["meta"] == {"toneladas": 3000, "falta": 2968, "percentual": 1.1, "dias_restantes": 13, "ritmo_necessario": 228.3}
    # O seguro entra no custo por tonelada, mas nao no lucro real.
    assert (r["despesas"]["empresa"], r["lucro_real"], r["sobra"]) == (800, 800, 300)
    assert r["precificacao"] == {"custo_fixo": 1600, "por_tonelada_atual": 50, "por_tonelada_na_meta": 0.53,
                                 "ponto_de_equilibrio_ton": 32.0}
    assert [p["valor"] for p in r["cascata"]] == [8640, -6720, -320, 0, 1600, -800, 800, -500, 300]


# --------------------------------------------------------------------------
# Agenda de pagamentos
# --------------------------------------------------------------------------


def despesa(db, **campos):
    base = dict(escopo="empresa", grupo="Boletos", descricao="Acordo Buonny", dia_vencimento=15, valor=1312,
                competencia_inicio="2026-09")
    base.update(campos)
    d = Despesa(**base)
    db.add(d)
    db.flush()
    return d


def test_parcela_anda_sozinha_e_some_depois_da_ultima(db):
    despesa(db, parcela_inicial=5, parcelas_total=6)
    assert [i["parcela"] for i in fin.agenda(db, "2026-09")["itens"]] == ["5/6"]
    assert [i["parcela"] for i in fin.agenda(db, "2026-10")["itens"]] == ["6/6"]
    assert fin.agenda(db, "2026-11")["itens"] == []
    assert fin.agenda(db, "2026-08")["itens"] == []  # antes de comecar


def test_vencimento_no_dia_31_cai_no_ultimo_dia_do_mes(db):
    despesa(db, dia_vencimento=31)
    assert fin.agenda(db, "2026-09")["itens"][0]["vencimento"] == "2026-09-30"


def test_situacao_de_cada_conta(db):
    despesa(db, descricao="Vencida", dia_vencimento=10)
    despesa(db, descricao="Hoje", dia_vencimento=16)
    despesa(db, descricao="Semana que vem", dia_vencimento=20)
    despesa(db, descricao="Sem dia", dia_vencimento=None)
    despesa(db, descricao="So precificacao", conta_no_resultado=False)
    ag = fin.agenda(db, "2026-09", hoje=date(2026, 9, 16))
    assert {i["descricao"]: i["situacao"] for i in ag["itens"]} == {
        "Vencida": "atrasado", "Hoje": "hoje", "Semana que vem": "a_vencer", "Sem dia": "sem_data",
    }
    assert (ag["totais"]["atrasados"], ag["totais"]["proximos_7_dias"]) == (1, 2624)


def test_pagar_pelo_banco_lanca_a_saida_e_desfazer_tira(db):
    nu = conta(db, saldo=5000, em=date(2026, 9, 1))
    d = despesa(db, parcela_inicial=5, parcelas_total=6)
    fin.pagar(db, origem="despesa", origem_id=d.id, competencia="2026-09", pago_em=DIA, conta_id=nu.id)
    item = fin.agenda(db, "2026-09")["itens"][0]
    assert item["situacao"] == "pago" and item["pagamento"]["conta"] == "Nubank"
    lanc = db.query(LancamentoCaixa).one()
    assert (lanc.tipo, lanc.valor, lanc.descricao, lanc.origem) == ("saida", 1312, "Acordo Buonny 5/6", "agenda")
    with pytest.raises(fin.ErroFinanceiro):
        fin.pagar(db, origem="despesa", origem_id=d.id, competencia="2026-09", pago_em=DIA)
    fin.desfazer_pagamento(db, origem="despesa", origem_id=d.id, competencia="2026-09")
    assert db.query(LancamentoCaixa).count() == 0
    assert fin.agenda(db, "2026-09")["itens"][0]["situacao"] != "pago"


# --------------------------------------------------------------------------
# Importar as planilhas
# --------------------------------------------------------------------------


def planilha_fluxo() -> bytes:
    livro = Workbook()
    painel = livro.active
    painel.title = "Dashboard"
    painel["D1"], painel["E1"] = "Data:", datetime(2026, 9, 15)
    painel["A5"] = "Intituições Bancarias"
    for linha, (nome, aba) in enumerate((("Nu Pagamentos", "NU"), ("Banco do Brasil", "Banco do Brasil"), ("Banco Itau", "Itau")), start=6):
        painel.cell(linha, 1, nome)
        painel.cell(linha, 2, f"='{aba}'!$D$1" if " " in aba else f"={aba}!D1")
    movimentos = {
        "NU": (3128.22, [("TRNSF: Banco do Brasil", 10000, "PIX: Bianca Maria ", 20), ("BOLETOS: Equilibrio", 12672, "TRNSF: Banco Itau", 3000)]),
        "Banco do Brasil": (10632.08, [("PIX: Fertimaxi", 21824, "TRNSF: Nu Pagamentos", 10000)]),
        "Itau": (412.16, [("TRNSF: Nu Pagamentos", 3000, "DEBITO: Parcela Giro", 769.94), ("RENDE FACIL", 0.03, None, 0)]),
    }
    for aba, (saldo, linhas) in movimentos.items():
        folha = livro.create_sheet(aba)
        folha["B1"], folha["D1"] = "Saldo Anterior:", saldo
        folha["A4"], folha["B4"], folha["C4"], folha["D4"] = "Descrição", "Entradas", "Descrição ", "Saidas"
        for n, (desc_e, val_e, desc_s, val_s) in enumerate(linhas, start=5):
            folha.cell(n, 1, desc_e), folha.cell(n, 2, val_e), folha.cell(n, 3, desc_s), folha.cell(n, 4, val_s)
        folha.cell(5 + len(linhas), 2, 0)
        folha.cell(10, 1, "Total da receita")
    arquivo = io.BytesIO()
    livro.save(arquivo)
    return arquivo.getvalue()


def planilha_controle() -> bytes:
    livro = Workbook()
    lucro = livro.active
    lucro.title = "LUCRO BRUTO"
    lucro["D4"] = 3000
    linhas = [
        (5003, datetime(2026, 9, 1), "EDILSON TEIXEIRA", "FERTIMAXI", "Jaíba MG", 32, 270, 210, None, 10, 0, "QUEIROZ"),
        (5004, datetime(2026, 9, 1), "KAIFFER NATAN PEREIRA", "HERINGER", "Ervália MG", 32, 340, "-", 9000, 10, 0, "QUEIROZ"),
        (5010, " ", "CANCELADO", None, None, None, None, None, None, None, None, None),
    ]
    for n, (cte, dia, mot, fab, dest, peso, fe, fm, fm_total, ag, co, contr) in enumerate(linhas, start=6):
        for col, valor in ((1, cte), (2, dia), (3, mot), (4, fab), (5, dest), (6, peso), (7, fe), (9, fm), (11, ag), (13, co), (17, contr)):
            lucro.cell(n, col, valor)
        lucro.cell(n, 8, f"=G{n}*F{n}")
        lucro.cell(n, 10, fm_total if fm_total is not None else f"=I{n}*F{n}")
        lucro.cell(n, 12, f"=K{n}*F{n}")
        lucro.cell(n, 14, f"=M{n}*F{n}")

    def aba_despesas(nome, grupos):
        folha = livro.create_sheet(nome)
        folha["A4"], folha["B4"], folha["C4"] = "Despesas", "Vencimento", "Valor Despesa"
        n = 5
        for grupo, itens in grupos:
            if grupo:
                folha.cell(n, 1, grupo)
                n += 1
            for desc, venc, valor in itens:
                folha.cell(n, 1, desc), folha.cell(n, 2, venc), folha.cell(n, 3, valor)
                n += 1
            n += 1
        folha.cell(n + 2, 1, "Valor Despesas")

    aba_despesas("GASTOS EMPRESA", [
        ("SISTEMAS", [("BSOFT", "16", 614.21), ("GOFLUX", "16", 583.34)]),
        ("BOLETOS & ACORDOS ", [("ACORDO BUONNY 5/6", 15, 1312), ("ACORDO BUONNY 6/6", 30, 1312)]),
        ("CONSORCIOS & SEGUROS", [("ITAU CONSORCIO VEICULO ", 5, 1376.37), ("INTERNET - SALA", "-", 0)]),
    ])
    aba_despesas("GASTOS PESSOAIS", [("FAMILIA", [("ESCOLA", 5, 950)])])
    aba_despesas("PRECIFICAÇÃO CTE", [("SISTEMAS", [("BSOFT", "16", 614.21), ("SEGURO RCTRC", None, 4000)])])
    dividas = livro.create_sheet("DIVIDAS ATIVAS")
    for n, linha in enumerate((("RODRIGO", 70000, "3/7", 10000), ("HUGO", "A ORGANIZAR", None, 10000)), start=5):
        for col, valor in enumerate(linha, start=1):
            dividas.cell(n, col, valor)
    livro.create_sheet("PAGAMENTOS")["A2"] = "JULHO"
    arquivo = io.BytesIO()
    livro.save(arquivo)
    return arquivo.getvalue()


def test_fluxo_de_caixa_previa_nao_grava_e_aplicar_bate_o_saldo(db):
    previa = imp.importar_planilha(db, planilha_fluxo(), "Fluxo Caixa_15.09.xlsx")
    assert previa["tipo"] == "fluxo_caixa" and previa["aplicado"] is False
    assert db.query(ContaBancaria).count() == 0 and db.query(LancamentoCaixa).count() == 0

    feito = imp.importar_planilha(db, planilha_fluxo(), "Fluxo Caixa_15.09.xlsx", aplicar=True)
    assert feito["data"] == "2026-09-15"
    assert [(c["nome"], c["criada"]) for c in feito["contas"]] == [("Nubank", True), ("Banco do Brasil", True), ("Itaú", True)]
    assert feito["lancamentos"] == {"novos": 9, "ja_existiam": 0}
    # BB -> Nu (10.000) e Nu -> Itau (3.000) viram transferencias.
    assert feito["transferencias_ligadas"] == 2
    resumo = fin.resumo_caixa(db, DIA, DIA)
    assert {c["nome"]: c["fechamento"] for c in resumo["contas"]} == {
        "Nubank": 22780.22, "Banco do Brasil": 22456.08, "Itaú": 2642.25,
    }
    assert resumo["totais"]["transferencias"] == 13000
    rende = db.query(LancamentoCaixa).filter(LancamentoCaixa.forma == "RENDIMENTO").one()
    assert rende.descricao == "Rende Facil"


def test_fluxo_de_caixa_importado_de_novo_nao_duplica(db):
    imp.importar_planilha(db, planilha_fluxo(), "x.xlsx", aplicar=True)
    de_novo = imp.importar_planilha(db, planilha_fluxo(), "x.xlsx", aplicar=True)
    assert de_novo["lancamentos"] == {"novos": 0, "ja_existiam": 9}
    assert all(c["divergencia"] is None and not c["criada"] for c in de_novo["contas"])
    assert db.query(LancamentoCaixa).count() == 9


def test_controle_de_carregamentos_completo(db):
    feito = imp.importar_planilha(db, planilha_controle(), "Controle de carregamentos_Setembro.xlsx", aplicar=True)
    assert feito["competencia"] == "2026-09"
    assert feito["carregamentos"] == {"novos": 3, "ja_existiam": 0, "cancelados": 1}
    assert feito["despesas_empresa"] == {"novas": 6, "ja_existiam": 0, "total": 5197.92}
    assert feito["precificacao"] == {"marcadas": 1, "novas": 1}
    assert feito["dividas"] == {"novas": 2, "ja_existiam": 0}
    assert any("PAGAMENTOS" in aviso for aviso in feito["avisos"])

    r = fin.resultado_mensal(db, "2026-09", hoje=date(2026, 9, 16))
    # 5003: 32x(270-210-10)=1.600; 5004: 32x340 - 9.000 - 320 = 1.560
    assert (r["resumo"]["toneladas"], r["resumo"]["lucro_bruto"], r["meta"]["toneladas"]) == (64, 3160, 3000)
    assert (r["despesas"]["empresa"], r["despesas"]["pessoal"], r["precificacao"]["custo_fixo"]) == (5197.92, 950, 4614.21)

    buonny = db.query(Despesa).filter(Despesa.descricao == "Acordo Buonny").order_by(Despesa.parcela_inicial).all()
    assert [(d.parcela_inicial, d.parcelas_total, d.dia_vencimento) for d in buonny] == [(5, 6, 15), (6, 6, 30)]
    assert db.query(Despesa).filter(Despesa.descricao == "Itaú Consórcio Veículo").one().grupo == "Consórcios & Seguros"
    rodrigo, hugo = db.query(Divida).order_by(Divida.id).all()
    assert (rodrigo.parcelas_pagas, rodrigo.parcelas_total, rodrigo.valor_total) == (3, 7, 70000)
    assert (hugo.valor_total, hugo.observacao) == (None, "A organizar")

    de_novo = imp.importar_planilha(db, planilha_controle(), "Controle de carregamentos_Setembro.xlsx", aplicar=True)
    assert de_novo["carregamentos"]["novos"] == 0 and de_novo["despesas_empresa"]["novas"] == 0 and de_novo["dividas"]["novas"] == 0


def test_arquivo_que_nao_e_planilha_do_financeiro(db):
    livro = Workbook()
    arquivo = io.BytesIO()
    livro.save(arquivo)
    with pytest.raises(fin.ErroFinanceiro, match="não parece"):
        imp.importar_planilha(db, arquivo.getvalue(), "outra.xlsx")
    with pytest.raises(fin.ErroFinanceiro, match="abrir"):
        imp.importar_planilha(db, b"nao sou xlsx", "x.xlsx")


# --------------------------------------------------------------------------
# Extrato OFX
# --------------------------------------------------------------------------

OFX = b"""OFXHEADER:100
DATA:OFXSGML
CHARSET:1252

<OFX><BANKMSGSRSV1><STMTTRNRS><STMTRS><BANKTRANLIST>
<STMTTRN><TRNTYPE>CREDIT<DTPOSTED>20260915120000[-3:BRT]<TRNAMT>21824.00<FITID>A1<MEMO>Pix recebido - FERTIMAXI</STMTTRN>
<STMTTRN><TRNTYPE>DEBIT<DTPOSTED>20260915<TRNAMT>-45.90<FITID>A2<MEMO>TARIFA PACOTE SERVICOS</STMTTRN>
<STMTTRN><TRNTYPE>DEBIT<DTPOSTED>20260916<TRNAMT>-1221,00<FITID>A3<MEMO>PAGAMENTO DE TITULO</STMTTRN>
</BANKTRANLIST><LEDGERBAL><BALAMT>32090.10<DTASOF>20260916</LEDGERBAL></STMTRS></STMTTRNRS></BANKMSGSRSV1></OFX>
"""


def test_extrato_ofx_entra_sem_duplicar_o_que_ja_foi_lancado(db):
    bb = conta(db, "Banco do Brasil", saldo=10632.08, em=DIA)
    lancar(db, bb, DIA, "entrada", 21824, "Fertimaxi")  # ja veio da planilha
    db.commit()
    previa = imp.importar_extrato(db, bb.id, OFX)
    assert previa["lancamentos"] == {"novos": 2, "ja_existiam": 0, "parecidos_ignorados": 1}
    assert db.query(LancamentoCaixa).count() == 1
    assert [p["forma"] for p in previa["previa"]] == ["DEBITO", "BOLETO"]

    feito = imp.importar_extrato(db, bb.id, OFX, aplicar=True)
    assert feito["conferencia"] == {"data": "2026-09-16", "banco": 32090.10, "sistema": 31189.18, "diferenca": 900.92}
    de_novo = imp.importar_extrato(db, bb.id, OFX, aplicar=True)
    assert de_novo["lancamentos"]["novos"] == 0 and db.query(LancamentoCaixa).count() == 3


def test_arquivo_que_nao_e_ofx(db):
    bb = conta(db)
    with pytest.raises(fin.ErroFinanceiro, match="OFX"):
        imp.importar_extrato(db, bb.id, b"%PDF-1.4 nada aqui")


# --------------------------------------------------------------------------
# Rotas
# --------------------------------------------------------------------------


def cliente_http(db, papel="admin"):
    app = FastAPI()
    app.include_router(rotas.router)
    usuario = SimpleNamespace(email="dono@atlantico", role=papel)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: usuario
    return TestClient(app)


def test_somar_mes_fica_no_ultimo_dia_quando_o_mes_e_curto():
    assert fin.somar_mes(date(2026, 1, 31)) == date(2026, 2, 28)
    assert fin.somar_mes(date(2026, 12, 10)) == date(2027, 1, 10)


def test_parcela_paga_empurra_a_data_e_quitar_limpa(db):
    http = cliente_http(db)
    van = http.post("/financeiro/dividas", json={"credor": "Van", "valor_total": 30000, "parcelas_total": 10, "parcelas_pagas": 8,
                                                 "valor_parcela": 3000, "proximo_pagamento": "2026-09-30"}).json()
    assert van["proximo_pagamento"] == "2026-09-30"
    assert http.post(f"/financeiro/dividas/{van['id']}/parcela-paga").json()["proximo_pagamento"] == "2026-10-30"
    ultima = http.post(f"/financeiro/dividas/{van['id']}/parcela-paga").json()
    assert ultima["quitada"] is True and ultima["proximo_pagamento"] is None


def test_dividas_em_aberto_na_ordem_do_pagamento(db):
    http = cliente_http(db)
    for credor, dia in (("Sem data", None), ("Depois", "2026-10-25"), ("Antes", "2026-10-10")):
        http.post("/financeiro/dividas", json={"credor": credor, "proximo_pagamento": dia})
    assert [d["credor"] for d in http.get("/financeiro/dividas").json()] == ["Antes", "Depois", "Sem data"]


def test_so_administrador_ve_o_financeiro(db):
    assert cliente_http(db, papel="user").get("/financeiro/contas").status_code == 403
    assert cliente_http(db).get("/financeiro/contas").status_code == 200


def test_rotas_do_dia_a_dia(db):
    http = cliente_http(db)
    nu = http.post("/financeiro/contas", json={"nome": "Nubank", "saldo_inicial": 1000, "saldo_inicial_em": "2026-09-01"}).json()
    bb = http.post("/financeiro/contas", json={"nome": "BB", "saldo_inicial": 0, "saldo_inicial_em": "2026-09-01"}).json()
    assert http.post("/financeiro/lancamentos", json={"tipo": "entrada", "data": "2026-09-15", "valor": 250, "conta_id": nu["id"], "forma": "pix"}).status_code == 200
    assert http.post("/financeiro/lancamentos", json={"tipo": "transferencia", "data": "2026-09-15", "valor": 100,
                                                      "conta_id": nu["id"], "conta_destino_id": bb["id"]}).status_code == 200
    ruim = http.post("/financeiro/lancamentos", json={"tipo": "transferencia", "data": "2026-09-15", "valor": 100, "conta_id": nu["id"]})
    assert ruim.status_code == 400
    caixa = http.get("/financeiro/caixa", params={"inicio": "2026-09-15"}).json()
    assert caixa["totais"]["fechamento"] == 1250 and caixa["totais"]["transferencias"] == 100

    despesa_id = http.post("/financeiro/despesas", json={"escopo": "empresa", "descricao": "Bsoft", "dia_vencimento": 10, "valor": 614.21,
                                                         "competencia_inicio": "2026-09"}).json()["id"]
    pago = http.post("/financeiro/agenda/pagar", json={"origem": "despesa", "origem_id": despesa_id, "competencia": "2026-09",
                                                       "pago_em": "2026-09-10", "conta_id": nu["id"]})
    assert pago.status_code == 200 and pago.json()["totais"]["pago"] == 614.21
    duplicado = http.post("/financeiro/agenda/pagar", json={"origem": "despesa", "origem_id": despesa_id, "competencia": "2026-09", "pago_em": "2026-09-10"})
    assert duplicado.status_code == 400
    assert http.get("/financeiro/contas").json()[0]["saldo_atual"] == round(1250 - 100 - 614.21 + 100, 2) - 100

    http.post("/financeiro/despesas", json={"escopo": "pessoal", "descricao": "Escola", "dia_vencimento": 5, "valor": 950, "competencia_inicio": "2026-09"})
    vencidas = http.post("/financeiro/agenda/pagar-vencidas", json={"competencia": "2026-09", "ate": "2026-09-16"}).json()
    assert vencidas["marcados"] >= 1 and all(i["situacao"] == "pago" for i in vencidas["itens"] if i["vencimento"] <= "2026-09-16")

    divida = http.post("/financeiro/dividas", json={"credor": "Van", "valor_total": 15000, "parcelas_total": 10, "parcelas_pagas": 9, "valor_parcela": 3000}).json()
    assert divida["restante"] == 3000
    quitada = http.post(f"/financeiro/dividas/{divida['id']}/parcela-paga").json()
    assert quitada["quitada"] is True and quitada["restante"] == 0
    assert http.post(f"/financeiro/dividas/{divida['id']}/parcela-paga").status_code == 400

    importada = http.post("/financeiro/importar-planilha", files={"arquivo": ("Fluxo Caixa_15.09.xlsx", planilha_fluxo())})
    assert importada.status_code == 200 and importada.json()["aplicado"] is False
