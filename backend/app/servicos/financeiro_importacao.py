"""Traz pro sistema as planilhas do financeiro e o extrato do banco (OFX).

Tudo roda primeiro como previa: faz a importacao numa transacao, mostra o
que entraria e desfaz. So grava com `aplicar=True`. Importar o mesmo
arquivo de novo nao duplica nada.
"""
from __future__ import annotations

import io
import re
import warnings
from collections import Counter
from datetime import date, datetime, timedelta
from typing import Optional

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from ..models import CarregamentoFinanceiro, ContaBancaria, Despesa, Divida, LancamentoCaixa, MetaMensal
from . import financeiro as fin

MESES = {
    "janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5, "junho": 6, "julho": 7,
    "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12,
}

# Identidade de cada banco na tela (cor da marca).
BANCOS_CONHECIDOS = (
    ("nu pagamentos", "Nubank", "#820AD1"),
    ("nubank", "Nubank", "#820AD1"),
    ("c6", "C6 Bank", "#3A3A3A"),
    ("banco do brasil", "Banco do Brasil", "#F2C200"),
    ("sicredi", "Sicredi", "#3FA110"),
    ("itau", "Itaú", "#EC7000"),
)

_SIGLAS = {"BB", "C6", "NU", "PIX", "CDT", "FGTS", "RCTRC", "RCDC", "MSC", "GR", "PIC", "TED", "DOC", "CT", "CH",
           "MG", "BA", "RS", "SP", "RJ", "PR", "SC", "GO", "ES", "MAP", "EMP", "CNPJ", "CPF", "IPVA", "INSS", "RF"}
_ACENTOS = {
    "ITAU": "Itaú", "CARTAO": "Cartão", "CARTOES": "Cartões", "SALARIO": "Salário", "SALARIOS": "Salários",
    "CONSORCIO": "Consórcio", "CONSORCIOS": "Consórcios", "VEICULO": "Veículo", "ESCRITORIO": "Escritório",
    "CONVENIO": "Convênio", "DIVIDAS": "Dívidas", "FAMILIA": "Família", "ARVORES": "Árvores", "SAUDE": "Saúde",
    "AGUA": "Água", "RENEGOCIACAO": "Renegociação", "PRECIFICACAO": "Precificação", "MANUTENCAO": "Manutenção",
    "COMISSAO": "Comissão", "ONIBUS": "Ônibus", "GRAFICA": "Gráfica", "CONTABIL": "Contábil",
}
_MINUSCULAS = {"DE", "DA", "DO", "DAS", "DOS", "E"}


def nome_legivel(texto) -> str:
    """ "ITAU CONSORCIO VEICULO" -> "Itaú Consórcio Veículo" (sem cara de planilha)."""
    partes = re.findall(r"\w+|\W+", re.sub(r"\s+", " ", str(texto or "")).strip())
    saida, primeira = [], True
    for parte in partes:
        if not re.match(r"\w", parte):
            saida.append(parte)
            continue
        maiuscula = parte.upper()
        if maiuscula in _ACENTOS:
            saida.append(_ACENTOS[maiuscula])
        elif maiuscula in _SIGLAS or re.search(r"\d", parte):
            saida.append(maiuscula)
        elif maiuscula in _MINUSCULAS and not primeira:
            saida.append(parte.lower())
        else:
            saida.append(parte[:1].upper() + parte[1:].lower())
        primeira = False
    return "".join(saida).strip()


def _numero(valor) -> Optional[float]:
    if isinstance(valor, bool) or valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace("R$", "").strip()
    if re.fullmatch(r"-?[\d.]+,\d+", texto):
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def _dia(valor) -> Optional[int]:
    numero = _numero(valor)
    return int(numero) if numero and 1 <= numero <= 31 else None


def _abrir(conteudo: bytes):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        formulas = load_workbook(io.BytesIO(conteudo), data_only=False)
        valores = load_workbook(io.BytesIO(conteudo), data_only=True)
    return formulas, valores


def _aba(livro, *nomes):
    por_nome = {fin.sem_acento(n): n for n in livro.sheetnames}
    for nome in nomes:
        if fin.sem_acento(nome) in por_nome:
            return livro[por_nome[fin.sem_acento(nome)]]
    return None


def detectar(livro) -> Optional[str]:
    if _aba(livro, "LUCRO BRUTO") is not None:
        return "controle"
    if _aba(livro, "Dashboard") is not None and any(
        str((livro[n]["A4"].value or "")).strip().lower().startswith("descri") for n in livro.sheetnames
    ):
        return "fluxo_caixa"
    return None


def importar_planilha(db: Session, conteudo: bytes, nome_arquivo: str = "", *, aplicar: bool = False,
                      usuario: str = "", competencia: Optional[str] = None) -> dict:
    try:
        formulas, valores = _abrir(conteudo)
    except Exception as exc:
        raise fin.ErroFinanceiro("Não consegui abrir o arquivo: envie a planilha .xlsx") from exc
    tipo = detectar(valores)
    if tipo is None:
        raise fin.ErroFinanceiro(
            "Essa planilha não parece nem o Fluxo de Caixa nem o Controle de Carregamentos"
        )
    try:
        if tipo == "fluxo_caixa":
            resumo = _importar_fluxo(db, formulas, valores, nome_arquivo, usuario)
        else:
            resumo = _importar_controle(db, formulas, valores, nome_arquivo, competencia)
        if aplicar:
            db.commit()
        else:
            db.rollback()
    except Exception:
        db.rollback()
        raise
    return {"tipo": tipo, "aplicado": aplicar, **resumo}


# --------------------------------------------------------------------------
# Fluxo de caixa
# --------------------------------------------------------------------------


def _data_do_fluxo(valores, nome_arquivo: str) -> date:
    painel = _aba(valores, "Dashboard")
    for linha in painel.iter_rows(min_row=1, max_row=4):
        for celula in linha:
            if isinstance(celula.value, datetime):
                return celula.value.date()
    achou = re.search(r"(\d{1,2})[._-](\d{1,2})(?:[._-](\d{2,4}))?", nome_arquivo or "")
    if achou:
        ano = int(achou.group(3)) if achou.group(3) else date.today().year
        return date(ano + 2000 if ano < 100 else ano, int(achou.group(2)), int(achou.group(1)))
    raise fin.ErroFinanceiro("Não achei a data da planilha (célula 'Data:' do Dashboard)")


def _abas_de_banco(formulas, valores) -> list[tuple[str, str]]:
    """(nome do banco no painel, aba) - pela formula "=NU!D1" de cada linha."""
    painel_f, painel_v = _aba(formulas, "Dashboard"), _aba(valores, "Dashboard")
    achados = []
    for linha in range(1, 40):
        nome = painel_v.cell(linha, 1).value
        formula = painel_f.cell(linha, 2).value
        if not nome or not isinstance(formula, str):
            continue
        ref = re.match(r"^=\s*'?([^'!]+)'?!", formula)
        if ref and ref.group(1) in valores.sheetnames:
            achados.append((str(nome).strip(), ref.group(1)))
    return achados


def _conta_do_banco(db: Session, nome_no_painel: str, dia: date, saldo_anterior: float, ordem: int) -> tuple[ContaBancaria, dict]:
    chave = fin.sem_acento(nome_no_painel)
    apelido, cor = nome_legivel(nome_no_painel), "#64746C"
    for trecho, nome, cor_marca in BANCOS_CONHECIDOS:
        if trecho in chave:
            apelido, cor = nome, cor_marca
            break
    conta = next(
        (c for c in db.query(ContaBancaria).all()
         if fin.sem_acento(c.instituicao) == chave or fin.sem_acento(c.nome) == fin.sem_acento(apelido)),
        None,
    )
    if conta is None:
        conta = ContaBancaria(nome=apelido, instituicao=nome_legivel(nome_no_painel), cor=cor, saldo_inicial=fin.dinheiro(saldo_anterior),
                              saldo_inicial_em=dia, ordem=ordem)
        db.add(conta)
        db.flush()
        return conta, {"nome": conta.nome, "criada": True, "saldo_inicial": conta.saldo_inicial, "divergencia": None}
    no_sistema = fin.saldo_no_inicio(db, conta, dia)
    divergencia = None
    if abs(no_sistema - saldo_anterior) > 0.009:
        divergencia = {"sistema": no_sistema, "planilha": fin.dinheiro(saldo_anterior)}
    return conta, {"nome": conta.nome, "criada": False, "saldo_inicial": no_sistema, "divergencia": divergencia}


def _importar_fluxo(db: Session, formulas, valores, nome_arquivo: str, usuario: str) -> dict:
    dia = _data_do_fluxo(valores, nome_arquivo)
    bancos = _abas_de_banco(formulas, valores)
    if not bancos:
        raise fin.ErroFinanceiro("Não achei os bancos no Dashboard da planilha")

    contas, novos, repetidos = [], 0, 0
    saldo_planilha = 0.0
    for ordem, (nome, aba) in enumerate(bancos):
        folha = valores[aba]
        saldo_anterior = _numero(folha["D1"].value) or 0.0
        saldo_atual = _numero(folha["D2"].value)
        conta, info = _conta_do_banco(db, nome, dia, saldo_anterior, ordem)
        existentes = {
            l.id_externo for l in db.query(LancamentoCaixa.id_externo).filter(LancamentoCaixa.conta_id == conta.id)
        }
        for linha in range(5, folha.max_row + 1):
            rotulo = str(folha.cell(linha, 1).value or "")
            if fin.sem_acento(rotulo).startswith("total"):
                break
            for coluna_desc, coluna_valor, tipo, lado in ((1, 2, "entrada", "E"), (3, 4, "saida", "S")):
                valor = _numero(folha.cell(linha, coluna_valor).value)
                if not valor or valor <= 0:
                    continue
                id_externo = f"planilha:{dia.isoformat()}:{linha}:{lado}"
                if id_externo in existentes:
                    repetidos += 1
                    continue
                forma, descricao = fin.separar_forma(folha.cell(linha, coluna_desc).value)
                if descricao.isupper():
                    descricao = nome_legivel(descricao)  # "RENDE FACIL" -> "Rende Facil"
                fin.criar_lancamento(
                    db, conta_id=conta.id, data=dia, tipo=tipo, valor=valor, forma=forma,
                    descricao=descricao or "Sem descrição", usuario=usuario, origem="planilha", id_externo=id_externo,
                )
                novos += 1
        info["saldo_final_planilha"] = fin.dinheiro(saldo_atual) if saldo_atual is not None else None
        contas.append((conta, info))
        saldo_planilha += saldo_atual or 0.0

    ligadas = fin.ligar_transferencias(db, dia)
    resumo = fin.resumo_caixa(db, dia, dia, dias_serie=1)
    fechamento = {c["id"]: c["fechamento"] for c in resumo["contas"]}
    for conta, info in contas:
        info["saldo_final_sistema"] = fechamento.get(conta.id)
    return {
        "data": dia.isoformat(),
        "contas": [info for _, info in contas],
        "lancamentos": {"novos": novos, "ja_existiam": repetidos},
        "transferencias_ligadas": ligadas,
        "conferencia": {
            "saldo_final_planilha": fin.dinheiro(saldo_planilha),
            "saldo_final_sistema": fin.dinheiro(sum(fechamento.get(c.id, 0) for c, _ in contas)),
        },
    }


# --------------------------------------------------------------------------
# Controle de carregamentos
# --------------------------------------------------------------------------


def _competencia_do_controle(valores, nome_arquivo: str, informada: Optional[str]) -> str:
    if informada:
        return fin.validar_competencia(informada)
    folha = _aba(valores, "LUCRO BRUTO")
    meses = Counter(
        fin.competencia_de(c.value.date()) for c in folha["B"][5:] if isinstance(c.value, datetime)
    )
    if meses:
        return meses.most_common(1)[0][0]
    nome = fin.sem_acento(nome_arquivo)
    for mes, numero in MESES.items():
        if mes in nome:
            return f"{date.today().year:04d}-{numero:02d}"
    raise fin.ErroFinanceiro("Não achei o mês da planilha: informe a competência")


def _parte(folha_f, folha_v, linha: int, coluna_ton: int, coluna_total: int) -> tuple[Optional[float], Optional[float]]:
    """Por tonelada, ou total fechado quando a celula do total foi digitada
    (e nao a formula peso x valor)."""
    por_ton = _numero(folha_v.cell(linha, coluna_ton).value)
    formula_total = folha_f.cell(linha, coluna_total).value
    total = None
    if not (isinstance(formula_total, str) and formula_total.startswith("=")):
        total = _numero(formula_total)
    return por_ton, total


def _importar_controle(db: Session, formulas, valores, nome_arquivo: str, competencia: Optional[str]) -> dict:
    competencia = _competencia_do_controle(valores, nome_arquivo, competencia)
    avisos = []
    resumo = {"competencia": competencia}

    # --- LUCRO BRUTO
    folha_f, folha_v = _aba(formulas, "LUCRO BRUTO"), _aba(valores, "LUCRO BRUTO")
    existentes = {
        (fin.sem_acento(c.ctes), fin.sem_acento(c.motorista))
        for c in db.query(CarregamentoFinanceiro).filter(CarregamentoFinanceiro.competencia == competencia)
    }
    novos = repetidos = cancelados = 0
    for linha in range(6, folha_v.max_row + 1):
        cte = folha_v.cell(linha, 1).value
        motorista = str(folha_v.cell(linha, 3).value or "").strip()
        if cte in (None, "") and not motorista:
            continue
        ctes = str(int(cte) if isinstance(cte, float) and cte.is_integer() else cte or "").strip()
        cancelado = fin.sem_acento(motorista) == "cancelado"
        chave = (fin.sem_acento(ctes), "" if cancelado else fin.sem_acento(motorista))
        if chave in existentes:
            repetidos += 1
            continue
        emissao = folha_v.cell(linha, 2).value
        partes = {
            "frete_empresa": _parte(folha_f, folha_v, linha, 7, 8),
            "frete_motorista": _parte(folha_f, folha_v, linha, 9, 10),
            "agenciamento": _parte(folha_f, folha_v, linha, 11, 12),
            "comissao": _parte(folha_f, folha_v, linha, 13, 14),
        }
        carregamento = CarregamentoFinanceiro(
            competencia=competencia,
            ctes=ctes,
            data_emissao=emissao.date() if isinstance(emissao, datetime) else None,
            motorista="" if cancelado else motorista.upper(),
            fabrica=nome_legivel(folha_v.cell(linha, 4).value),
            destino=str(folha_v.cell(linha, 5).value or "").strip(),
            peso=_numero(folha_v.cell(linha, 6).value) or 0,
            contratante=nome_legivel(folha_v.cell(linha, 17).value),
            cancelado=cancelado,
            **{f"{p}_ton": v[0] for p, v in partes.items()},
            **{f"{p}_total": v[1] for p, v in partes.items()},
        )
        db.add(carregamento)
        existentes.add(chave)
        novos += 1
        cancelados += int(cancelado)
    db.flush()
    resumo["carregamentos"] = {"novos": novos, "ja_existiam": repetidos, "cancelados": cancelados}

    meta_planilha = _numero(folha_v["D4"].value)
    if meta_planilha:
        meta = db.get(MetaMensal, competencia)
        if meta is None:
            db.add(MetaMensal(competencia=competencia, meta_toneladas=meta_planilha))
            db.flush()
        elif abs(float(meta.meta_toneladas or 0) - meta_planilha) > 0.001:
            avisos.append(f"A meta do mês já estava em {meta.meta_toneladas:g} t no sistema; a da planilha ({meta_planilha:g} t) não foi usada.")
    resumo["meta_toneladas"] = meta_planilha

    # --- GASTOS EMPRESA / PESSOAIS
    for aba, escopo in (("GASTOS EMPRESA", "empresa"), ("GASTOS PESSOAIS", "pessoal")):
        folha = _aba(valores, aba)
        resumo[f"despesas_{escopo}"] = _importar_despesas(db, folha, escopo, competencia) if folha is not None else None

    # --- PRECIFICACAO CTE
    folha = _aba(valores, "PRECIFICACAO CTE", "PRECIFICAÇÃO CTE")
    if folha is not None:
        resumo["precificacao"] = _importar_precificacao(db, folha, competencia)

    # --- DIVIDAS ATIVAS
    folha = _aba(valores, "DIVIDAS ATIVAS", "DÍVIDAS ATIVAS")
    if folha is not None:
        resumo["dividas"] = _importar_dividas(db, folha)

    if _aba(valores, "PAGAMENTOS") is not None:
        avisos.append(
            "A aba PAGAMENTOS não foi importada: ela estava com o calendário de outro mês (título JULHO, "
            "dia 1 num sábado e dia 31). No sistema a agenda se monta sozinha pelos vencimentos das "
            "despesas; posto, cheques e acertos entram como conta avulsa."
        )

    conferido = fin.resultado_mensal(db, competencia)
    resumo["conferencia"] = {
        "toneladas": {"planilha": _numero(folha_v["F4"].value), "sistema": conferido["resumo"]["toneladas"]},
        "lucro_bruto": {"planilha": _numero(folha_v["O4"].value), "sistema": conferido["resumo"]["lucro_bruto"]},
        "despesas_empresa": {"planilha": _total_da_aba(valores, "GASTOS EMPRESA"), "sistema": conferido["despesas"]["empresa"]},
        "gastos_pessoais": {"planilha": _total_da_aba(valores, "GASTOS PESSOAIS"), "sistema": conferido["despesas"]["pessoal"]},
        "custo_fixo_precificacao": {"planilha": _total_da_aba(valores, "PRECIFICACAO CTE"), "sistema": conferido["precificacao"]["custo_fixo"]},
    }
    resumo["avisos"] = avisos
    return resumo


def _total_da_aba(valores, aba: str) -> Optional[float]:
    folha = _aba(valores, aba)
    if folha is None:
        return None
    for linha in range(1, folha.max_row + 1):
        if fin.sem_acento(folha.cell(linha, 1).value).startswith("valor despesas"):
            return _numero(folha.cell(linha, 3).value)
    return None


def _linhas_de_despesa(folha):
    """(grupo, descricao, dia, valor) de cada linha, ate 'Valor Despesas'."""
    grupo = ""
    for linha in range(5, folha.max_row + 1):
        a, b, c = (folha.cell(linha, col).value for col in (1, 2, 3))
        texto = str(a or "").strip()
        if fin.sem_acento(texto).startswith("valor despesas"):
            break
        if not texto:
            continue
        if b is None and c is None:
            grupo = nome_legivel(texto)
            continue
        yield grupo, texto, _dia(b), _numero(c)


def _separar_parcela(descricao: str) -> tuple[str, Optional[int], Optional[int]]:
    achou = re.search(r"\s(\d{1,3})\s*/\s*(\d{1,3})\s*$", descricao)
    if not achou:
        return descricao, None, None
    return descricao[: achou.start()].strip(), int(achou.group(1)), int(achou.group(2))


def _chave_despesa(escopo, descricao, parcela):
    return (escopo, fin.sem_acento(descricao), parcela)


def _importar_despesas(db: Session, folha, escopo: str, competencia: str) -> dict:
    existentes = {_chave_despesa(d.escopo, d.descricao, d.parcela_inicial) for d in db.query(Despesa).filter(Despesa.escopo == escopo)}
    novas = repetidas = 0
    total = 0.0
    for ordem, (grupo, texto, dia, valor) in enumerate(_linhas_de_despesa(folha)):
        descricao, parcela, parcelas = _separar_parcela(texto)
        descricao = nome_legivel(descricao)
        total += valor or 0
        chave = _chave_despesa(escopo, descricao, parcela)
        if chave in existentes:
            repetidas += 1
            continue
        db.add(Despesa(
            escopo=escopo, grupo=grupo or "Outras", descricao=descricao, dia_vencimento=dia, valor=fin.dinheiro(valor),
            parcela_inicial=parcela, parcelas_total=parcelas, competencia_inicio=competencia, ordem=ordem,
        ))
        existentes.add(chave)
        novas += 1
    db.flush()
    return {"novas": novas, "ja_existiam": repetidas, "total": fin.dinheiro(total)}


def _importar_precificacao(db: Session, folha, competencia: str) -> dict:
    empresa = db.query(Despesa).filter(Despesa.escopo == "empresa").all()
    marcadas = novas = 0
    for ordem, (_, texto, dia, valor) in enumerate(_linhas_de_despesa(folha)):
        descricao, parcela, parcelas = _separar_parcela(texto)
        descricao = nome_legivel(descricao)
        existente = next((d for d in empresa if fin.sem_acento(d.descricao) == fin.sem_acento(descricao)), None)
        if existente is not None:
            if not existente.entra_precificacao:
                existente.entra_precificacao = True
                marcadas += 1
            continue
        # Custo que so existe na precificacao (seguro, Buonny): nao mexe no
        # lucro real do mes, que a planilha tambem nao descontava.
        nova = Despesa(
            escopo="empresa", grupo="Custos da operação", descricao=descricao, dia_vencimento=dia,
            valor=fin.dinheiro(valor), parcela_inicial=parcela, parcelas_total=parcelas, competencia_inicio=competencia,
            conta_no_resultado=False, entra_precificacao=True, ordem=900 + ordem,
        )
        db.add(nova)
        empresa.append(nova)
        novas += 1
    db.flush()
    return {"marcadas": marcadas, "novas": novas}


def _importar_dividas(db: Session, folha) -> dict:
    existentes = {
        (fin.sem_acento(d.credor), fin.dinheiro(d.valor_total) if d.valor_total is not None else None,
         fin.dinheiro(d.valor_parcela) if d.valor_parcela is not None else None)
        for d in db.query(Divida).all()
    }
    novas = repetidas = 0
    for linha in range(5, folha.max_row + 1):
        credor = str(folha.cell(linha, 1).value or "").strip()
        if not credor or fin.sem_acento(credor).startswith("valor despesas"):
            if fin.sem_acento(credor).startswith("valor despesas"):
                break
            continue
        total = _numero(folha.cell(linha, 2).value)
        parcela = _numero(folha.cell(linha, 4).value)
        parcela = parcela if parcela else None
        chave = (fin.sem_acento(nome_legivel(credor)), fin.dinheiro(total) if total is not None else None,
                 fin.dinheiro(parcela) if parcela is not None else None)
        if chave in existentes:
            repetidas += 1
            continue
        pagas, quantidade = 0, None
        achou = re.fullmatch(r"\s*(\d+)\s*/\s*(\d+)\s*", str(folha.cell(linha, 3).value or ""))
        if achou:
            pagas, quantidade = int(achou.group(1)), int(achou.group(2))
        observacao = "A organizar" if "organizar" in fin.sem_acento(folha.cell(linha, 2).value) + fin.sem_acento(folha.cell(linha, 3).value) else ""
        db.add(Divida(credor=nome_legivel(credor), valor_total=total, parcelas_total=quantidade, parcelas_pagas=pagas,
                      valor_parcela=parcela, observacao=observacao))
        existentes.add(chave)
        novas += 1
    db.flush()
    return {"novas": novas, "ja_existiam": repetidas}


# --------------------------------------------------------------------------
# Extrato do banco (OFX)
# --------------------------------------------------------------------------


def _campo_ofx(bloco: str, nome: str) -> str:
    achou = re.search(rf"<{nome}>([^<\r\n]*)", bloco, flags=re.IGNORECASE)
    return achou.group(1).strip() if achou else ""


def _data_ofx(texto: str) -> Optional[date]:
    achou = re.match(r"(\d{4})(\d{2})(\d{2})", texto or "")
    return date(int(achou.group(1)), int(achou.group(2)), int(achou.group(3))) if achou else None


def ler_ofx(conteudo: bytes) -> dict:
    texto = None
    for codificacao in ("utf-8", "cp1252", "latin-1"):
        try:
            texto = conteudo.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    if not texto or "<STMTTRN>" not in texto.upper():
        raise fin.ErroFinanceiro("Esse arquivo não parece um extrato OFX (não tem lançamentos)")
    transacoes = []
    for indice, bloco in enumerate(re.findall(r"<STMTTRN>(.*?)</STMTTRN>", texto, flags=re.IGNORECASE | re.DOTALL)):
        valor = _numero(_campo_ofx(bloco, "TRNAMT").replace(",", "."))
        dia = _data_ofx(_campo_ofx(bloco, "DTPOSTED"))
        if valor in (None, 0) or dia is None:
            continue
        memo = _campo_ofx(bloco, "MEMO") or _campo_ofx(bloco, "NAME")
        transacoes.append({
            "fitid": _campo_ofx(bloco, "FITID") or f"{dia.isoformat()}:{valor}:{memo}:{indice}",
            "data": dia, "valor": abs(valor), "tipo": "entrada" if valor > 0 else "saida", "descricao": memo,
        })
    saldo = None
    bloco_saldo = re.search(r"<LEDGERBAL>(.*?)</LEDGERBAL>", texto, flags=re.IGNORECASE | re.DOTALL)
    if bloco_saldo:
        valor = _numero(_campo_ofx(bloco_saldo.group(1), "BALAMT").replace(",", "."))
        dia = _data_ofx(_campo_ofx(bloco_saldo.group(1), "DTASOF"))
        if valor is not None:
            saldo = {"valor": fin.dinheiro(valor), "data": dia}
    return {"transacoes": transacoes, "saldo": saldo}


def importar_extrato(db: Session, conta_id: int, conteudo: bytes, *, aplicar: bool = False, usuario: str = "") -> dict:
    conta = db.get(ContaBancaria, conta_id)
    if conta is None:
        raise fin.ErroFinanceiro("Conta bancária não encontrada")
    lido = ler_ofx(conteudo)
    try:
        existentes = {l.id_externo for l in db.query(LancamentoCaixa.id_externo).filter(LancamentoCaixa.conta_id == conta.id)}
        ja_lancados = db.query(LancamentoCaixa).filter(LancamentoCaixa.conta_id == conta.id, LancamentoCaixa.id_externo.is_(None) | ~LancamentoCaixa.id_externo.like("ofx:%")).all()
        # O que ja foi digitado ou veio da planilha (mesmo dia, tipo e valor)
        # nao entra de novo pelo extrato.
        parecidos = Counter((l.data, l.tipo, fin.dinheiro(l.valor)) for l in ja_lancados)
        novos, repetidos, parecidos_ignorados, previa = 0, 0, 0, []
        for t in lido["transacoes"]:
            id_externo = f"ofx:{t['fitid']}"[:160]
            if id_externo in existentes:
                repetidos += 1
                continue
            chave = (t["data"], t["tipo"], fin.dinheiro(t["valor"]))
            if parecidos[chave] > 0:
                parecidos[chave] -= 1
                parecidos_ignorados += 1
                continue
            forma = fin.forma_pelo_texto(t["descricao"])
            lanc = fin.criar_lancamento(
                db, conta_id=conta.id, data=t["data"], tipo=t["tipo"], valor=t["valor"], forma=forma,
                descricao=t["descricao"] or "Sem descrição", usuario=usuario, origem="extrato", id_externo=id_externo,
            )
            existentes.add(id_externo)
            novos += 1
            if len(previa) < 60:
                previa.append({"data": lanc.data.isoformat(), "tipo": lanc.tipo, "forma": lanc.forma,
                               "descricao": lanc.descricao, "valor": fin.dinheiro(lanc.valor)})
        for dia in sorted({t["data"] for t in lido["transacoes"]}):
            fin.ligar_transferencias(db, dia)

        conferencia = None
        if lido["saldo"] and lido["saldo"]["data"]:
            no_sistema = fin.saldo_no_inicio(db, conta, lido["saldo"]["data"] + timedelta(days=1))
            conferencia = {
                "data": lido["saldo"]["data"].isoformat(), "banco": lido["saldo"]["valor"], "sistema": no_sistema,
                "diferenca": fin.dinheiro(lido["saldo"]["valor"] - no_sistema),
            }
        datas = [t["data"] for t in lido["transacoes"]]
        resumo = {
            "tipo": "extrato", "aplicado": aplicar, "conta": conta.nome,
            "periodo": {"inicio": min(datas).isoformat(), "fim": max(datas).isoformat()} if datas else None,
            "lancamentos": {"novos": novos, "ja_existiam": repetidos, "parecidos_ignorados": parecidos_ignorados},
            "previa": previa, "conferencia": conferencia,
        }
        if aplicar:
            db.commit()
        else:
            db.rollback()
        return resumo
    except Exception:
        db.rollback()
        raise
